from openpyxl import Workbook
from openpyxl.styles import (
    Font,
    PatternFill,
    Border,
    Side,
    Alignment
)
from openpyxl.utils import get_column_letter


class ExcelExporter:

    def __init__(self):

        # ==================================================
        # Styles
        # ==================================================

        self.header_fill = PatternFill(
            fill_type="solid",
            start_color="1F4E78",
            end_color="1F4E78"
        )

        self.header_font = Font(
            bold=True,
            color="FFFFFF"
        )

        self.section_fill = PatternFill(
            fill_type="solid",
            start_color="D9EAF7",
            end_color="D9EAF7"
        )

        self.section_font = Font(
            bold=True
        )

        self.thin_border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin")
        )

    # ==================================================
    # Style Sheet
    # ==================================================

    def style_sheet(self, ws):

        # --------------------------------------------------
        # First row
        # --------------------------------------------------

        for cell in ws[1]:

            cell.fill = self.header_fill

            cell.font = self.header_font

            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True
            )

            cell.border = self.thin_border

        # --------------------------------------------------
        # Remaining rows
        # --------------------------------------------------

        for row in ws.iter_rows(
            min_row=2
        ):

            for cell in row:

                cell.border = self.thin_border

                cell.alignment = Alignment(
                    vertical="top",
                    wrap_text=True
                )

        # --------------------------------------------------
        # Auto width
        # --------------------------------------------------

        for column in ws.columns:

            max_length = 0

            column_letter = get_column_letter(
                column[0].column
            )

            for cell in column:

                try:

                    if cell.value is not None:

                        value_length = len(
                            str(cell.value)
                        )

                        max_length = max(
                            max_length,
                            value_length
                        )

                except Exception:

                    pass

            ws.column_dimensions[
                column_letter
            ].width = min(
                max_length + 3,
                60
            )

        # --------------------------------------------------
        # Freeze header
        # --------------------------------------------------

        ws.freeze_panes = "A2"

        # --------------------------------------------------
        # Filter
        # --------------------------------------------------

        if ws.max_row >= 1:

            ws.auto_filter.ref = ws.dimensions

    # ==================================================
    # Convert Value For Excel
    # ==================================================

    @staticmethod
    def excel_value(value):

        if value is None:

            return ""

        # --------------------------------------------------
        # Lists
        # --------------------------------------------------

        if isinstance(
            value,
            list
        ):

            return "\n".join(
                str(item)
                for item in value
                if item is not None
            )

        # --------------------------------------------------
        # Dictionaries
        # --------------------------------------------------

        if isinstance(
            value,
            dict
        ):

            lines = []

            for key, item in value.items():

                if isinstance(
                    item,
                    list
                ):

                    item = ", ".join(
                        str(x)
                        for x in item
                    )

                lines.append(
                    f"{key}: {item}"
                )

            return "\n".join(
                lines
            )

        return value

    # ==================================================
    # Header Sheet
    # ==================================================

    def create_header_sheet(
        self,
        wb,
        header
    ):

        ws = wb.active

        ws.title = "Header"

        ws.append([
            "Field",
            "Value"
        ])

        if not isinstance(
            header,
            dict
        ):

            header = {}

        for key, value in header.items():

            display_name = (
                key.replace(
                    "_",
                    " "
                )
                .title()
            )

            ws.append([
                display_name,
                self.excel_value(
                    value
                )
            ])

        self.style_sheet(
            ws
        )

        # Make value column wider
        ws.column_dimensions[
            "B"
        ].width = 60

        return ws

    # ==================================================
    # Container Sheet
    # ==================================================

    def create_container_sheet(
        self,
        wb,
        containers_data
    ):

        ws = wb.create_sheet(
            "Containers"
        )

        ws.append([
            "Container Number",
            "Seal Number",
            "Size",
            "Cartons",
            "Weight (KG)",
            "CBM"
        ])

        # --------------------------------------------------
        # Get containers
        # --------------------------------------------------

        containers = []

        if isinstance(
            containers_data,
            dict
        ):

            containers = containers_data.get(
                "containers",
                []
            )

        if not isinstance(
            containers,
            list
        ):

            containers = []

        # --------------------------------------------------
        # Add rows
        # --------------------------------------------------

        for container in containers:

            if not isinstance(
                container,
                dict
            ):

                continue

            ws.append([
                self.excel_value(
                    container.get(
                        "container_number"
                    )
                ),

                self.excel_value(
                    container.get(
                        "seal_number"
                    )
                ),

                self.excel_value(
                    container.get(
                        "size"
                    )
                ),

                self.excel_value(
                    container.get(
                        "cartons"
                    )
                ),

                self.excel_value(
                    container.get(
                        "weight_kg"
                    )
                ),

                self.excel_value(
                    container.get(
                        "cbm"
                    )
                )
            ])

        self.style_sheet(
            ws
        )

        # --------------------------------------------------
        # Widths
        # --------------------------------------------------

        ws.column_dimensions[
            "A"
        ].width = 22

        ws.column_dimensions[
            "B"
        ].width = 20

        ws.column_dimensions[
            "C"
        ].width = 15

        ws.column_dimensions[
            "D"
        ].width = 15

        ws.column_dimensions[
            "E"
        ].width = 18

        ws.column_dimensions[
            "F"
        ].width = 15

        return ws

    # ==================================================
    # Summary Sheet
    # ==================================================

    def create_summary_sheet(
        self,
        wb,
        summary
    ):

        ws = wb.create_sheet(
            "Summary"
        )

        ws.append([
            "Field",
            "Value"
        ])

        if not isinstance(
            summary,
            dict
        ):

            summary = {}

        # --------------------------------------------------
        # Keep the four summary fields in fixed order
        # --------------------------------------------------

        fields = [

            "total_containers",

            "total_cartons",

            "total_weight",

            "total_cbm",

        ]

        for field in fields:

            value = summary.get(
                field
            )

            ws.append([
                field.replace(
                    "_",
                    " "
                ).title(),

                self.excel_value(
                    value
                )
            ])

        self.style_sheet(
            ws
        )

        ws.column_dimensions[
            "A"
        ].width = 25

        ws.column_dimensions[
            "B"
        ].width = 25

        return ws

    # ==================================================
    # Validation Sheet
    # ==================================================

    def create_validation_sheet(
        self,
        wb,
        containers_data
    ):

        ws = wb.create_sheet(
            "Validation"
        )

        ws.append([
            "Validation",
            "Status",
            "Calculated",
            "Document"
        ])

        validation = {}

        if isinstance(
            containers_data,
            dict
        ):

            validation = containers_data.get(
                "validation",
                {}
            )

        if not isinstance(
            validation,
            dict
        ):

            validation = {}

        # --------------------------------------------------
        # Overall status
        # --------------------------------------------------

        overall_status = validation.get(
            "status"
        )

        ws.append([
            "Overall Status",
            self.excel_value(
                overall_status
            ),
            "",
            ""
        ])

        # --------------------------------------------------
        # Container count
        # --------------------------------------------------

        self._append_validation_row(
            ws,
            "Container Count",
            validation.get(
                "container_count"
            )
        )

        # --------------------------------------------------
        # Cartons
        # --------------------------------------------------

        self._append_validation_row(
            ws,
            "Cartons",
            validation.get(
                "cartons"
            )
        )

        # --------------------------------------------------
        # Weight
        # --------------------------------------------------

        self._append_validation_row(
            ws,
            "Weight (KG)",
            validation.get(
                "weight_kg"
            )
        )

        # --------------------------------------------------
        # CBM
        # --------------------------------------------------

        self._append_validation_row(
            ws,
            "CBM",
            validation.get(
                "cbm"
            )
        )

        # --------------------------------------------------
        # Duplicate containers
        # --------------------------------------------------

        duplicates = validation.get(
            "duplicate_container_numbers",
            []
        )

        ws.append([
            "Duplicate Container Numbers",

            self.excel_value(
                duplicates
            ),

            "",
            ""
        ])

        # --------------------------------------------------
        # Invalid container numbers
        # --------------------------------------------------

        invalid_numbers = validation.get(
            "invalid_container_numbers",
            []
        )

        ws.append([
            "Invalid Container Numbers",

            self.excel_value(
                invalid_numbers
            ),

            "",
            ""
        ])

        # --------------------------------------------------
        # Missing fields
        # --------------------------------------------------

        missing_fields = validation.get(
            "missing_fields",
            []
        )

        ws.append([
            "Missing Fields",

            self.excel_value(
                missing_fields
            ),

            "",
            ""
        ])

        # --------------------------------------------------
        # Calculated totals
        # --------------------------------------------------

        calculated_totals = validation.get(
            "calculated_totals",
            {}
        )

        if not isinstance(
            calculated_totals,
            dict
        ):

            calculated_totals = {}

        ws.append([
            "Calculated Total Containers",
            "",
            self.excel_value(
                calculated_totals.get(
                    "total_containers"
                )
            ),
            ""
        ])

        ws.append([
            "Calculated Total Cartons",
            "",
            self.excel_value(
                calculated_totals.get(
                    "total_cartons"
                )
            ),
            ""
        ])

        ws.append([
            "Calculated Total Weight",
            "",
            self.excel_value(
                calculated_totals.get(
                    "total_weight"
                )
            ),
            ""
        ])

        ws.append([
            "Calculated Total CBM",
            "",
            self.excel_value(
                calculated_totals.get(
                    "total_cbm"
                )
            ),
            ""
        ])

        # --------------------------------------------------
        # Style
        # --------------------------------------------------

        self.style_sheet(
            ws
        )

        ws.column_dimensions[
            "A"
        ].width = 35

        ws.column_dimensions[
            "B"
        ].width = 25

        ws.column_dimensions[
            "C"
        ].width = 20

        ws.column_dimensions[
            "D"
        ].width = 20

        return ws

    # ==================================================
    # Validation Row Helper
    # ==================================================

    def _append_validation_row(
        self,
        ws,
        name,
        validation_data
    ):

        if not isinstance(
            validation_data,
            dict
        ):

            validation_data = {}

        ws.append([
            name,

            self.excel_value(
                validation_data.get(
                    "status"
                )
            ),

            self.excel_value(
                validation_data.get(
                    "calculated"
                )
            ),

            self.excel_value(
                validation_data.get(
                    "document"
                )
            )
        ])

    # ==================================================
    # Main Export
    # ==================================================

    def export(
        self,
        data,
        output_path
    ):
        """
        Create the complete Excel workbook.

        Sheets:

        1. Header
        2. Containers
        3. Summary
        4. Validation
        """

        wb = Workbook()

        # ==================================================
        # Safety
        # ==================================================

        if not isinstance(
            data,
            dict
        ):

            data = {}

        header = data.get(
            "header",
            {}
        )

        containers = data.get(
            "containers",
            {}
        )

        summary = data.get(
            "summary",
            {}
        )

        # ==================================================
        # Header
        # ==================================================

        self.create_header_sheet(
            wb,
            header
        )

        # ==================================================
        # Containers
        # ==================================================

        self.create_container_sheet(
            wb,
            containers
        )

        # ==================================================
        # Summary
        # ==================================================

        self.create_summary_sheet(
            wb,
            summary
        )

        # ==================================================
        # Validation
        # ==================================================

        self.create_validation_sheet(
            wb,
            containers
        )

        # ==================================================
        # Save
        # ==================================================

        wb.save(
            output_path
        )

        print(
            "\nExcel workbook created:"
        )

        print(
            output_path
        )